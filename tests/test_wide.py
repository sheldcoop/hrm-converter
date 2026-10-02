"""Reference limits and the per-feature wide workbook."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from conftest import ConfigFactory
from create_test_fixture import FixtureInfo, synthetic_value
from hrm_converter.config import Config
from hrm_converter.models import LONG_COLUMNS, OutputError, Severity
from hrm_converter.pipeline import run_conversion
from hrm_converter.reference import ReferenceFileError, load_reference, parse_reference
from hrm_converter.validation import IssueCollector
from hrm_converter.wide import (
    WideError,
    WideResult,
    build_wide,
    main,
    read_long,
    write_reference_template,
    write_wide,
)

META = {
    "Project_Name": "SYN_Project",
    "Part_Number": "SYN-PART",
    "Lot_Number": "90001",
    "Lot_Name": "",
    "Process": "Post AAA",
    "Panel": 2,
    "Location": "Unit",
    "Source_File": "synthetic_Summary.xlsx",
}


def long_table(rows: list[tuple[str, str, object, str, int, str, object, str]]) -> pd.DataFrame:
    """rows: (Buildup, Side, Unit, Feature_Type, Feature_Number, Metric, Value, unit of measure)."""
    records = [
        {**META, "Buildup": b, "Side": s, "Unit": u, "Feature_Type": ft, "Feature_Number": fn,
         "Metric": m, "Value": v, "Unit_of_Measurement": uom}
        for b, s, u, ft, fn, m, v, uom in rows
    ]  # fmt: skip
    return pd.DataFrame(records, columns=list(LONG_COLUMNS), dtype=object)


LONG = long_table(
    [
        ("BU-01", "Front", 1, "Pad", 1, "Radius", 42.0, "µm"),
        ("BU-01", "Front", 1, "Pad", 1, "Dimple", 0.0, "µm"),
        ("BU-01", "Front", 2, "Pad", 1, "Radius", 44.0, "µm"),
        ("BU-01", "Front", 1, "Pad", 2, "Radius", 44.0, "µm"),
        ("BU-03", "Back", 1, "Pad", 1, "Radius", 42.0, "µm"),
        ("BU-03", "Back", "C1", "Pad", 1, "Radius", 40.0, "µm"),
        ("BU-01", "Front", 1, "Trace", 1, "Width", 21.5, "µm"),
        ("BU-01", "Front", 1, "Trace", 1, "Distance 1", 3.0, ""),
        ("BU-01", "Front", 1, "Roughness", 1, "Ra_Mean", 107.0, "nm"),
        ("BU-01", "Front", 1, "Roughness", 1, "Rz_Mean", 727.0, "nm"),
    ]
)


def reference_frame(rows: list[dict[str, object]]) -> pd.DataFrame:
    columns = ["Part_Number", "Buildup", "Side", "Feature_Type", "Feature_Number", "Metric",
               "LSL", "Target", "USL"]  # fmt: skip
    return pd.DataFrame(rows, columns=columns, dtype=object)


def wide(
    rows: list[dict[str, object]], config: Config, long: pd.DataFrame = LONG
) -> tuple[WideResult, IssueCollector]:
    issues = IssueCollector()
    reference = parse_reference(reference_frame(rows), "ref.xlsx", config, issues)
    return build_wide(long, reference, config, issues), issues


def test_one_sheet_per_feature_with_one_column_per_metric(config: Config) -> None:
    result = build_wide(LONG, None, config, IssueCollector())
    assert list(result.sheets) == ["Pad", "Roughness", "Trace"]
    pad = result.sheets["Pad"]
    assert list(pad.columns) == [
        "Project_Name", "Part_Number", "Lot_Number", "Lot_Name", "Buildup", "Process", "Panel",
        "Side", "Location", "Unit", "Pad", "Radius (µm)", "Dimple (µm)", "Source_File",
    ]  # fmt: skip
    assert len(pad) == 5
    first = pad.iloc[0]
    assert (first["Buildup"], first["Unit"], first["Pad"]) == ("BU-01", 1, 1)
    assert (first["Radius (µm)"], first["Dimple (µm)"]) == (42.0, 0.0)  # zero is kept
    assert pd.isna(pad.iloc[1]["Dimple (µm)"])  # not measured stays blank
    # Metric order follows the config; unknown metrics come last, without a unit suffix.
    assert list(result.sheets["Trace"].columns)[-3:-1] == ["Width (µm)", "Distance 1"]
    assert list(result.sheets["Roughness"].columns)[-3:-1] == ["Rz_Mean (nm)", "Ra_Mean (nm)"]
    assert result.row_count == 7 and result.out_of_spec_count == 0
    assert "Out_of_Spec" not in pad.columns


def test_wide_rows_are_sorted_by_feature_number_then_unit(config: Config) -> None:
    pad = build_wide(LONG, None, config, IssueCollector()).sheets["Pad"]
    order = list(zip(pad["Buildup"], pad["Pad"], pad["Unit"], strict=True))
    assert order == [
        ("BU-01", 1, 1),
        ("BU-01", 1, 2),
        ("BU-01", 2, 1),
        ("BU-03", 1, 1),
        ("BU-03", 1, "C1"),
    ]


def test_blank_key_cells_apply_to_all(config: Config) -> None:
    result, issues = wide(
        [{"Feature_Type": "Pad", "Metric": "Radius", "LSL": 41, "USL": 43}], config
    )
    pad = result.sheets["Pad"]
    assert set(pad["Radius_LSL"]) == {41.0} and set(pad["Radius_USL"]) == {43.0}
    assert "Radius_Target" not in pad.columns  # no target anywhere: no empty column
    assert list(pad["Out_of_Spec"]) == ["", "Radius > USL", "Radius > USL", "", "Radius < LSL"]
    assert result.flagged["Pad"] == [(1, "Radius (µm)"), (2, "Radius (µm)"), (4, "Radius (µm)")]
    assert result.out_of_spec_count == 3
    assert issues.reportable() == []


def test_most_specific_row_wins(config: Config) -> None:
    rows: list[dict[str, object]] = [
        {"Feature_Type": "Pad", "Metric": "Radius", "LSL": 41, "USL": 43},
        {"Part_Number": "SYN-PART", "Buildup": "BU03", "Side": "back", "Feature_Type": "pad",
         "Metric": "radius", "LSL": 39, "USL": 41},
        {"Buildup": "BU-01", "Feature_Type": "Pad", "Feature_Number": 2, "Metric": "Radius",
         "USL": 45, "Target": 44},
    ]  # fmt: skip
    result, issues = wide(rows, config)
    pad = result.sheets["Pad"]
    limits = [(r["Buildup"], r["Pad"], r["Radius_LSL"], r["Radius_USL"]) for _, r in pad.iterrows()]
    assert limits[0] == ("BU-01", 1, 41.0, 43.0)
    assert limits[2][:2] == ("BU-01", 2) and pd.isna(limits[2][2]) and limits[2][3] == 45.0
    assert limits[3] == ("BU-03", 1, 39.0, 41.0)  # 'BU03' / 'back' / 'pad' still match
    assert list(pad["Out_of_Spec"]) == ["", "Radius > USL", "", "Radius > USL", ""]
    assert list(result.reference["Matched_Groups"]) == [1, 1, 1]  # type: ignore[index]
    assert issues.reportable() == []


def test_equally_specific_rows_with_different_limits_are_a_conflict(config: Config) -> None:
    rows: list[dict[str, object]] = [
        {"Buildup": "BU-01", "Feature_Type": "Pad", "Metric": "Radius", "USL": 43},
        {"Side": "Front", "Feature_Type": "Pad", "Metric": "Radius", "USL": 45},
    ]
    result, issues = wide(rows, config)
    pad = result.sheets["Pad"]
    # BU-01 Front fits both rows equally: no limit is applied and nothing is guessed.
    assert "Radius_USL" not in pad.columns and "Out_of_Spec" not in pad.columns
    assert [i.category for i in issues.reportable()] == ["reference_conflict"]  # reported once
    assert "rows 2, 3" in issues.reportable()[0].message


def test_bad_reference_rows_are_reported_and_ignored(config: Config) -> None:
    rows: list[dict[str, object]] = [
        {"Feature_Type": "Pad", "Metric": "Radius", "LSL": 44, "USL": 43},
        {"Feature_Type": "Pad", "Metric": "Dimple", "USL": "tbd"},
        {"Feature_Type": "Pad", "USL": 1},
        {"Feature_Type": "Pad", "Metric": "Bump"},
        {},
        {"Feature_Type": "Via", "Metric": "Dimple", "USL": "0,9"},
    ]
    result, issues = wide(rows, config)
    found = [(i.severity, i.category) for i in issues.reportable()]
    assert found == [
        (Severity.ERROR, "reference_row"),
        (Severity.ERROR, "reference_row"),
        (Severity.ERROR, "reference_row"),
        (Severity.WARNING, "reference_row"),
        (Severity.WARNING, "reference_unused"),
    ]
    messages = [i.message for i in issues.reportable()]
    assert "LSL (44) is greater than USL (43)" in messages[0]
    assert "'tbd'" in messages[1] and "row(s) 7" in messages[4]
    assert "Out_of_Spec" not in result.sheets["Pad"].columns


def test_reference_needs_the_required_columns(config: Config) -> None:
    frame = pd.DataFrame({"Feature": ["Pad"], "USL": [1]})
    with pytest.raises(ReferenceFileError, match="needs the columns Feature_Type and Metric"):
        parse_reference(frame, "ref.xlsx", config, IssueCollector())


def test_header_spelling_is_forgiving(config: Config) -> None:
    frame = pd.DataFrame(
        {"feature type": ["Pad"], "METRIC": ["Radius"], "usl": [43], "Note": ["x"]}
    )
    reference = parse_reference(frame, "ref.xlsx", config, IssueCollector())
    assert reference.key_columns == ("Feature_Type", "Metric")
    assert reference.rules[0].limits.usl == 43.0


def test_duplicate_measurements_become_repeat_rows(config: Config) -> None:
    long = long_table(
        [
            ("BU-01", "Front", 1, "Via", 1, "Dimple", 0.5, "µm"),
            ("BU-01", "Front", 1, "Via", 1, "Dimple", 0.7, "µm"),
        ]
    )
    via = build_wide(long, None, config, IssueCollector()).sheets["Via"]
    assert list(via["Dimple (µm)"]) == [0.5, 0.7] and list(via["Repeat"]) == [1, 2]


def test_long_table_without_required_columns_is_rejected(config: Config) -> None:
    with pytest.raises(WideError, match="lacks the column"):
        build_wide(LONG.drop(columns=["Metric"]), None, config, IssueCollector())


def test_wide_workbook_is_written_with_tables_and_highlights(
    config: Config, tmp_path: Path
) -> None:
    result, _ = wide([{"Feature_Type": "Pad", "Metric": "Radius", "LSL": 41, "USL": 43}], config)
    path = write_wide(result, tmp_path / "wide.xlsx")
    workbook = openpyxl.load_workbook(path)
    assert workbook.sheetnames == [
        "Pad",
        "Roughness",
        "Trace",
        "Limit_Summary",
        "Reference",
        "Wide_Issues",
    ]
    pad = workbook["Pad"]
    assert pad.freeze_panes == "A2" and "tblPad" in pad.tables
    header = [cell.value for cell in pad[1]]
    radius = header.index("Radius (µm)") + 1
    fills = [pad.cell(row=row, column=radius).fill.fgColor.rgb for row in range(2, 7)]
    assert [fill.endswith("FFC7CE") for fill in fills] == [False, True, True, False, True]
    summary = pd.read_excel(path, sheet_name="Limit_Summary")
    radius_row = summary[(summary["Feature_Type"] == "Pad") & (summary["Metric"] == "Radius")].iloc[
        0
    ]
    assert (radius_row["Values"], radius_row["With_Limits"], radius_row["Out_of_Spec"]) == (5, 5, 3)
    with pytest.raises(OutputError, match="must end with .xlsx"):
        write_wide(result, tmp_path / "wide.csv")


def test_end_to_end_from_folders_to_wide_workbook(
    hierarchy: FixtureInfo, make_config: ConfigFactory, tmp_path: Path
) -> None:
    config = make_config()
    long_path = run_conversion(hierarchy.lot_a1, config).output_path
    assert long_path is not None

    reference_path = tmp_path / "reference.xlsx"
    write_reference_template(read_long(long_path, config), reference_path, config)
    template = pd.read_excel(reference_path, sheet_name="Limits", dtype=object)
    assert {"Pad", "Trace", "Via", "Roughness", "Fiducial"} <= set(template["Feature_Type"])
    with pytest.raises(OutputError, match="was not overwritten"):
        write_reference_template(read_long(long_path, config), reference_path, config)

    # Fill in one limit, as a user would in Excel: Pad Radius on BU-01 front.
    row = template[(template["Feature_Type"] == "Pad") & (template["Metric"] == "Radius")].index[0]
    template.loc[row, ["Buildup", "Side", "LSL", "USL"]] = ["BU01", "Front", 100.0, 100.25]
    filled = tmp_path / "filled.xlsx"
    template.to_excel(filled, sheet_name="Limits", index=False)

    wide_path = tmp_path / "wide out" / "wide.xlsx"
    assert (
        main(["--input", str(long_path), "--reference", str(filled), "--output", str(wide_path)])
        == 0
    )
    book = pd.read_excel(wide_path, sheet_name=None, dtype=object, keep_default_na=False)
    assert {
        "Pad",
        "SR Pad",
        "Trace",
        "Via",
        "Roughness",
        "Fiducial",
        "Limit_Summary",
        "Reference",
    } <= set(book)
    pad = book["Pad"]
    limited = pad[pad["Radius_USL"] != ""]
    assert set(limited["Buildup"]) == {"BU-01"} and set(limited["Side"]) == {"Front"}
    assert set(limited["Location"]) == {"Unit"}  # coupon front has no Radius block here
    # Panel 2 front, Pad 1: seed 100, Radius is column 3 -> unit 3 gives 100.33 > USL 100.25.
    panel2 = limited[(limited["Panel"] == 2) & (limited["Pad"] == 1)]
    assert list(panel2["Radius (µm)"]) == [synthetic_value(100, u, 3) for u in (1, 2, 3)]
    assert list(panel2["Out_of_Spec"]) == ["", "", "Radius > USL"]
    loaded = load_reference(filled, config, IssueCollector())
    assert len(loaded.rules) == 1 and loaded.source == "filled.xlsx [Limits]"


def test_wide_cli_reports_a_missing_long_workbook(
    tmp_path: Path, capsys: pytest.CaptureFixture[str], monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.chdir(tmp_path)
    assert main(["--input", str(tmp_path / "missing.xlsx")]) == 1
    assert "Long workbook cannot be read" in capsys.readouterr().err
