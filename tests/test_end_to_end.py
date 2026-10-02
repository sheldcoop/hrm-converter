"""End-to-end runs on the synthetic hierarchy, checked through the output workbook."""

from __future__ import annotations

import hashlib
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from conftest import ConfigFactory
from create_test_fixture import (
    LOT_A1,
    PART_A,
    PART_B,
    PROJECT,
    FixtureInfo,
    make_block,
    synthetic_value,
    write_workbook,
)
from hrm_converter.cli import main
from hrm_converter.config import Config, build_config
from hrm_converter.models import (
    ISSUE_COLUMNS,
    LONG_COLUMNS,
    LONG_LIMIT_COLUMNS,
    LONG_SHEET_COLUMNS,
    SUMMARY_COLUMNS,
    OutputError,
)
from hrm_converter.pipeline import run_conversion

EXPECTED_PROJECT_ROWS = 298


def digests(info: FixtureInfo) -> dict[Path, tuple[str, int]]:
    return {
        path: (hashlib.sha256(path.read_bytes()).hexdigest(), path.stat().st_mtime_ns)
        for path in info.all_source_files
    }


def sheets(path: Path | None) -> dict[str, pd.DataFrame]:
    assert path is not None
    return pd.read_excel(path, sheet_name=None, dtype=object, keep_default_na=False)


def long_of(folder: Path, config: Config) -> pd.DataFrame:
    return sheets(run_conversion(folder, config).output_path)["Long"]


def test_project_run_writes_the_expected_workbook(hierarchy: FixtureInfo, config: Config) -> None:
    before = digests(hierarchy)
    result = run_conversion(hierarchy.project, config)
    assert digests(hierarchy) == before  # source files are never modified
    assert sorted(p for p in hierarchy.root.rglob("*") if p.is_file()) == hierarchy.all_source_files

    assert (result.hrm_folder_count, result.processed_count, result.skipped_count) == (5, 10, 4)
    assert len(result.records) == EXPECTED_PROJECT_ROWS

    book = sheets(result.output_path)
    assert list(book) == [
        "Long", "Processing_Summary", "Validation_Issues", "Folder_Check", "Fix_Plan",
    ]  # fmt: skip
    assert tuple(book["Long"].columns) == LONG_SHEET_COLUMNS
    assert all(set(book["Long"][c]) == {""} for c in LONG_LIMIT_COLUMNS)  # no reference given
    assert tuple(book["Processing_Summary"].columns) == SUMMARY_COLUMNS
    assert tuple(book["Validation_Issues"].columns) == ISSUE_COLUMNS

    long = book["Long"]
    assert len(long) == EXPECTED_PROJECT_ROWS
    assert set(long["Project_Name"]) == {PROJECT}
    assert set(long["Part_Number"]) == {PART_A, PART_B}
    assert set(long["Lot_Number"]) == {"90001", "90002", "80001"}
    assert list(dict.fromkeys(long["Buildup"])) == ["BU-01", "BU-02", "BU-10"]  # natural order
    assert set(long["Location"]) == {"Unit", "Coupon"}
    assert set(long["Feature_Type"]) == {
        "Pad", "SR Pad", "Trace", "Via", "Roughness", "Landing_Pad", "Fiducial",
    }  # fmt: skip
    assert {"Distance 1", "Distance 2", "Angle", "Area"} <= set(long["Metric"])
    assert not long["Source_File"].str.contains("/").any()
    assert not long["Value"].eq("").any()  # blank cells produce no rows


def test_expected_records(hierarchy: FixtureInfo, config: Config) -> None:
    long = long_of(hierarchy.project, config)

    def value(**where: object) -> object:
        mask = pd.Series(True, index=long.index)
        for column, wanted in where.items():
            mask &= long[column] == wanted
        rows = long[mask]
        assert len(rows) == 1, f"{where} matched {len(rows)} rows"
        return rows.iloc[0]["Value"]

    full = {"Lot_Number": LOT_A1, "Buildup": "BU-01", "Process": "Post AAA", "Panel": 2,
            "Side": "Front", "Location": "Unit"}  # fmt: skip
    # Pad 2, Unit 3, 4th column (Radius): seed 200.
    assert value(**full, Feature_Type="Pad", Feature_Number=2, Unit=3, Metric="Radius") == (
        pytest.approx(synthetic_value(200, 3, 3))
    )
    # Roughness: 1st column is Rz_Mean whatever the operator typed; µm -> nm.
    assert value(**full, Feature_Type="Roughness", Unit=1, Metric="Rz_Mean") == pytest.approx(
        synthetic_value(0, 1, 0) * 1000
    )
    assert value(**full, Feature_Type="Roughness", Unit=1, Metric="Ra_Max") == pytest.approx(
        synthetic_value(0, 1, 7) * 1000
    )
    units = long[(long["Feature_Type"] == "Roughness")]["Unit_of_Measurement"]
    assert set(units) == {"nm"}

    coupon = long[(long["Location"] == "Coupon") & (long["Side"] == "Front")]
    assert set(coupon["Unit"]) == {"C1", "C2", "C3"}
    assert set(coupon["Feature_Type"]) == {"Trace"}

    future = {"Lot_Number": LOT_A1, "Buildup": "BU-01", "Panel": 10, "Feature_Type": "Pad"}
    assert value(**future, Unit=1, Metric="Dimple") == 0  # zero is preserved
    assert value(**future, Unit=3, Metric="Max_Stepheight") == 12.5  # '12,5' stored as text
    pad = long[(long["Panel"] == 10) & (long["Feature_Type"] == "Pad")]
    assert len(pad) == 7  # blank Radius of unit 1 and text 'n/a' of unit 2 give no rows
    assert not ((pad["Unit"] == 1) & (pad["Metric"] == "Radius")).any()
    assert not ((pad["Unit"] == 2) & (pad["Metric"] == "Max_Stepheight")).any()

    special = long[long["Process"] == "Special Run"]
    assert set(special["Feature_Type"]) == {"Via"} and set(special["Feature_Number"]) == {1, 2}

    numeric = long["Value"].map(lambda v: isinstance(v, int | float) and not isinstance(v, bool))
    assert numeric.all()


def test_summary_and_issues_are_recorded(hierarchy: FixtureInfo, config: Config) -> None:
    book = sheets(run_conversion(hierarchy.project, config).output_path)
    summary = book["Processing_Summary"]
    assert len(summary) == 14
    assert summary["Status"].value_counts().to_dict() == {"Processed": 10, "Skipped": 4}
    assert summary["Row_Count"].sum() == EXPECTED_PROJECT_ROWS
    reasons = " | ".join(summary[summary["Status"] == "Skipped"]["Reason"])
    assert "Multiple workbooks" in reasons
    assert "Metadata conflict between hierarchy and workbook: Buildup, Side" in reasons
    assert "cannot be opened" in reasons

    issues = book["Validation_Issues"]
    by_category = issues.groupby("Category")["Severity"].agg(list).to_dict()
    assert by_category["multiple_workbooks"] == ["ERROR"]
    assert by_category["unreadable_workbook"] == ["ERROR"]
    assert by_category["text_value"] == ["ERROR"]
    assert by_category["no_workbook"] == ["WARNING"]
    assert by_category["missing_hrm_folder"] == ["WARNING"]
    assert by_category["new_feature_type"] == ["WARNING"]
    assert by_category["numeric_text"] == ["WARNING"]
    assert by_category["header_order"] == ["WARNING"]
    assert len(by_category["empty_folder"]) == 2
    assert set(by_category["metadata_conflict"]) == {"ERROR"}

    conflicts = issues[issues["Category"] == "metadata_conflict"]
    assert set(conflicts["Field"]) == {"Buildup", "Side"}
    buildup = conflicts[conflicts["Field"] == "Buildup"].iloc[0]
    assert (buildup["Hierarchy_Value"], buildup["Workbook_Value"]) == ("BU-01", "BU09")
    assert buildup["Sheet"] == "SYN_SYN-PART-A_BU09_Front"
    assert buildup["Source_File"].endswith("_BU09_Front_Summary.xlsx")


def test_output_is_reproducible(hierarchy: FixtureInfo, config: Config) -> None:
    first = sheets(run_conversion(hierarchy.project, config).output_path)
    second = sheets(run_conversion(hierarchy.project, config).output_path)
    for name in first:
        pd.testing.assert_frame_equal(first[name], second[name])


@pytest.mark.parametrize(
    ("level", "rows", "lots", "buildups"),
    [
        ("part_a", 274, {"90001", "90002"}, {"BU-01", "BU-02", "BU-10"}),
        ("lot_a1", 244, {"90001"}, {"BU-01", "BU-02", "BU-10"}),
        ("buildup_a1_bu01", 208, {"90001"}, {"BU-01"}),
    ],
)
def test_selection_scope(
    hierarchy: FixtureInfo,
    config: Config,
    level: str,
    rows: int,
    lots: set[str],
    buildups: set[str],
) -> None:
    long = long_of(getattr(hierarchy, level), config)
    assert len(long) == rows
    assert set(long["Lot_Number"]) == lots
    assert set(long["Buildup"]) == buildups
    # Parent metadata is reconstructed even when a Lot or Buildup is selected.
    assert set(long["Project_Name"]) == {PROJECT}
    assert set(long["Part_Number"]) == {PART_A}


def test_permissive_mode_uses_the_hierarchy_value(
    hierarchy: FixtureInfo, make_config: ConfigFactory
) -> None:
    config = make_config(processing={"strict_metadata_conflicts": False})
    result = run_conversion(hierarchy.buildup_a1_bu01, config)
    book = sheets(result.output_path)
    long = book["Long"]
    assert len(long) == 208 + 18
    conflict_rows = long[long["Source_File"].str.endswith("_BU09_Front_Summary.xlsx")]
    assert set(conflict_rows["Buildup"]) == {"BU-01"} and set(conflict_rows["Side"]) == {"Back"}
    issues = book["Validation_Issues"]
    assert set(issues[issues["Category"] == "metadata_conflict"]["Severity"]) == {"WARNING"}


def test_output_workbook_formatting(hierarchy: FixtureInfo, config: Config) -> None:
    result = run_conversion(hierarchy.lot_a1, config)
    assert result.output_path is not None
    workbook = openpyxl.load_workbook(result.output_path)
    for name, table in (
        ("Long", "tblLong"),
        ("Processing_Summary", "tblProcessingSummary"),
        ("Validation_Issues", "tblValidationIssues"),
    ):
        sheet = workbook[name]
        assert sheet.freeze_panes == "A2"
        assert table in sheet.tables
        assert sheet.tables[table].ref == sheet.dimensions
    long = workbook["Long"]
    value_column = LONG_COLUMNS.index("Value") + 1
    cells = [row[value_column - 1] for row in long.iter_rows(min_row=2)]
    assert all(cell.data_type == "n" for cell in cells)  # numeric, and no formulas
    assert long.column_dimensions["A"].width >= 10


def test_output_is_never_read_back(hierarchy: FixtureInfo, tmp_path: Path, config: Config) -> None:
    # A private hierarchy whose output is written next to the source folders.
    side = (
        tmp_path
        / "P"
        / "PN"
        / "L1"
        / "BU01"
        / "HRM"
        / "post AAA_HRM-2099-0001"
        / "Panel 1"
        / "front"
    )
    write_workbook(side / "a_Summary.xlsx", "S", [make_block("Pad 1", ("Radius",), 1)])
    inside = build_config(
        {
            "output": {"directory": str(side), "filename": "out_Summary.xlsx"},
            "logging": {"directory": str(tmp_path / "logs")},
        }
    )
    with pytest.raises(OutputError, match="lies inside a 'HRM' source folder"):
        run_conversion(tmp_path / "P", inside)

    beside = build_config(
        {
            "output": {"directory": str(tmp_path / "P" / "output")},
            "logging": {"directory": str(tmp_path / "logs")},
        }
    )
    first = run_conversion(tmp_path / "P", beside)
    second = run_conversion(tmp_path / "P", beside)
    assert len(first.records) == len(second.records) == 3
    assert second.processed_count == 1


def test_extra_sheets_and_real_layout_quirks(tmp_path: Path, config: Config) -> None:
    side = (
        tmp_path
        / "P"
        / "PN"
        / "L1"
        / "BU03"
        / "HRM"
        / "post DDV_HRM-2099-0001"
        / "Panel 05"
        / "back"
    )
    write_workbook(
        side / "2099-01-01_00-00-00_SYN_P_BU03_Back_Summary.xlsx",
        "SYN_PN_BU03_Back",
        [
            make_block(
                "SR Pad 2", ("Max Stepheight", "Radius"), 5, units=(7, 8, 9, 4, 5, 6, 1, 2, 3)
            )
        ],
        extra_sheets=["Notes"],
    )
    result = run_conversion(tmp_path / "P" / "PN" / "L1", config)
    book = sheets(result.output_path)
    long = book["Long"]
    assert len(long) == 18
    assert list(dict.fromkeys(long["Unit"])) == [1, 2, 3, 4, 5, 6, 7, 8, 9]
    assert set(long["Feature_Type"]) == {"SR Pad"} and set(long["Feature_Number"]) == {2}
    assert set(long["Buildup"]) == {"BU-03"} and set(long["Panel"]) == {5}
    assert set(long["Process"]) == {"Post DDV"}
    assert list(book["Validation_Issues"]["Category"]) == ["extra_sheets"]


def test_cli_run_prints_a_summary_without_measurements(
    hierarchy: FixtureInfo,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.chdir(tmp_path)  # default ./logs and ./config.yaml are relative to the cwd
    exit_code = main(
        [
            "--input",
            str(hierarchy.lot_a1),
            "--output-dir",
            str(tmp_path / "cli out"),
            "--lot-name",
            "DOE Lot",
        ]
    )
    output = capsys.readouterr().out
    assert exit_code == 0
    assert "Lot Number folder" in output
    assert "Workbooks processed : 8" in output
    assert "Output rows         : 244" in output
    assert "Warnings / errors" in output
    assert str(synthetic_value(100, 1, 0)) not in output
    (written,) = (tmp_path / "cli out").glob("HRM_Long_SYN-PART-A_90001_*.xlsx")
    long = sheets(written)["Long"]
    assert set(long["Lot_Name"]) == {"DOE Lot"}
    assert len(list((tmp_path / "logs").glob("hrm_converter_*.log"))) == 1


def test_cli_reports_a_bad_folder_clearly(
    tmp_path: Path, capsys: pytest.CaptureFixture[str], monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.chdir(tmp_path)
    (tmp_path / "empty").mkdir()
    assert main(["--input", str(tmp_path / "empty"), "--output-dir", str(tmp_path / "o")]) == 1
    assert "No 'HRM' folder was found" in capsys.readouterr().err


def test_open_file_link_sits_next_to_source_file(hierarchy: FixtureInfo, config: Config) -> None:
    from hrm_converter.output_writer import OPEN_FILE_LABEL, open_file_link

    result = run_conversion(hierarchy.buildup_a1_bu01, config)
    assert result.output_path is not None
    sheet = openpyxl.load_workbook(result.output_path)["Long"]
    header = [cell.value for cell in sheet[1]]
    source = header.index("Source_File")
    assert header[source + 1] == "Open_File"

    for row in sheet.iter_rows(min_row=2):
        formula = row[source + 1].value
        assert formula.startswith("=HYPERLINK(") and formula.endswith(f',"{OPEN_FILE_LABEL}")')
        target = "".join(formula[len("=HYPERLINK(") : formula.rindex(",")].split('"&"')).strip('"')
        assert Path(target).is_file() and Path(target).name == row[source].value
        assert row[source + 1].font.underline == "single"

    # Paths longer than Excel's 255-character literal limit are split, not cut.
    long_path = "L:\\" + "\\".join(["a folder with a long name"] * 14) + '\\it"s.xlsx'
    formula = open_file_link(long_path)
    assert formula is not None and len(long_path) > 255
    literals = formula[len("=HYPERLINK(") : formula.rindex(",")].split("&")
    assert all(len(literal) <= 202 for literal in literals)
    assert "".join(literal[1:-1] for literal in literals).replace('""', '"') == long_path
    assert open_file_link("") is None  # uploads have no path to open
