"""Read-only access to HRM summary workbooks.

An HRM summary workbook has one sheet with measurement blocks stacked on top
of each other::

            |          | Max Stepheight | Min Stepheight | Radius | ...   <- header row
    Unit 7  | Pad 1    | 16.62          | 16.49          | 42.25  |       <- first unit row
    Unit 8  |          | 17.56          | 17.47          | 42.37  |
            |          | Width          | Space          |                <- next header row
    Unit 7  | Trace 1  | 21.52          | 33.35          |

Column A holds the unit, column B the feature label (first row of a block
only) and the columns from C onward the metric values. Workbooks are opened
read-only and are never written to.
"""

from __future__ import annotations

import re
from collections.abc import Sequence
from pathlib import Path

import openpyxl

from hrm_converter.models import CellValue, RawBlock, RawRow, RawSheet, WorkbookReadError

_UNIT_LABEL = re.compile(r"^(?:unit|coupon|c)?[\s\-_.]*\d+$", re.IGNORECASE)


def _text(value: object) -> str | None:
    """Trimmed text of a cell, or None when the cell is blank."""
    if value is None:
        return None
    text = " ".join(str(value).split())
    return text or None


def _cell(value: object) -> CellValue:
    """Keep numbers as numbers; everything else becomes trimmed text or None."""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int | float):
        return value
    return _text(value)


def _is_unit_label(value: object) -> bool:
    if isinstance(value, bool):
        return False
    if isinstance(value, int):
        return True
    return isinstance(value, str) and bool(_UNIT_LABEL.match(value.strip()))


class _BlockBuilder:
    def __init__(self, start_row: int, label: str | None, header: tuple[str | None, ...]) -> None:
        self.start_row = start_row
        self.label = label
        self.header = header
        self.rows: list[RawRow] = []

    def units(self) -> set[str]:
        return {row.unit_label.lower() for row in self.rows}

    def build(self) -> RawBlock:
        return RawBlock(self.start_row, self.label, self.header, tuple(self.rows))


def parse_rows(
    rows: Sequence[Sequence[object]],
) -> tuple[tuple[RawBlock, ...], tuple[int, ...]]:
    """Split sheet rows into measurement blocks; returns (blocks, stray row numbers)."""
    builders: list[_BlockBuilder] = []
    stray: list[int] = []
    header: tuple[str | None, ...] = ()
    current: _BlockBuilder | None = None

    for number, raw in enumerate(rows, start=1):
        row = tuple(raw) + (None,) * max(0, 2 - len(raw))
        unit_cell, label, rest = row[0], _text(row[1]), row[2:]
        if _text(unit_cell) is None:
            if any(isinstance(v, str) and v.strip() for v in rest):
                cells = [_text(v) for v in rest]
                while cells and cells[-1] is None:
                    cells.pop()
                header = tuple(cells)
                current = None
            elif label is not None or any(_text(v) is not None for v in rest):
                stray.append(number)
            continue
        if not _is_unit_label(unit_cell):
            stray.append(number)
            continue

        unit_label = " ".join(str(unit_cell).split())
        note: str | None = None
        # A labelled row whose unit was already seen starts a new block that
        # reuses the previous header (the operator left out the header row).
        repeats = (
            current is not None and label is not None and unit_label.lower() in current.units()
        )
        if current is None or repeats:
            current = _BlockBuilder(number, label, header)
            builders.append(current)
        else:
            note = label
        current.rows.append(RawRow(number, unit_label, note, tuple(_cell(v) for v in rest)))

    return tuple(b.build() for b in builders), tuple(stray)


def read_workbook(path: Path) -> RawSheet:
    """Read the measurement blocks of the first sheet of an HRM workbook."""
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # corrupted, password-protected, not a real xlsx, locked ...
        raise WorkbookReadError(
            f"Workbook cannot be opened (corrupted, password-protected or not a valid "
            f".xlsx file): {type(exc).__name__}: {exc}"
        ) from exc
    try:
        names = workbook.sheetnames
        if not names:
            raise WorkbookReadError("Workbook contains no sheets.")
        sheet = workbook[names[0]]
        rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
    except WorkbookReadError:
        raise
    except Exception as exc:
        raise WorkbookReadError(
            f"Workbook content cannot be read: {type(exc).__name__}: {exc}"
        ) from exc
    finally:
        workbook.close()

    blocks, stray = parse_rows(rows)
    return RawSheet(names[0], blocks, stray, tuple(names[1:]))
