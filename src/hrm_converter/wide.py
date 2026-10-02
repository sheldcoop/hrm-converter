"""Wide format: one sheet per feature type, derived from the long table.

The long table stays the single source. This module only reshapes it:

* one sheet per ``Feature_Type`` (``Pad``, ``SR Pad``, ``Trace``, ...; a new
  feature type gets its own sheet automatically);
* one row per unit and feature number, one column per metric;
* ``<Metric>_LSL`` / ``_Target`` / ``_USL`` columns filled from the reference
  sheet (see ``reference.py``), and an ``Out_of_Spec`` column naming every
  metric outside its limits.

Run:
    python -m hrm_converter.wide --input output/hrm_long_format.xlsx --reference HRM_Reference.xlsx
"""

from __future__ import annotations

import argparse
import os
import re
import sys
from collections.abc import Sequence
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from hrm_converter.config import Config, load_config
from hrm_converter.logging_setup import close_logging, setup_logging
from hrm_converter.models import LONG_COLUMNS, HrmConverterError, Issue, OutputError
from hrm_converter.output_writer import format_sheet, issues_frame
from hrm_converter.reference import (
    LIMIT_COLUMNS,
    LimitMatcher,
    Limits,
    Reference,
    is_blank,
    load_reference,
    reference_template,
)
from hrm_converter.validation import IssueCollector, natural_key

META_COLUMNS: tuple[str, ...] = LONG_COLUMNS[:9]  # Project_Name .. Location
SUMMARY_SHEET = "Limit_Summary"
REFERENCE_SHEET = "Reference"
ISSUES_SHEET = "Wide_Issues"
SUMMARY_COLUMNS = ("Feature_Type", "Metric", "Values", "With_Limits", "Out_of_Spec")
OUT_OF_SPEC = "Out_of_Spec"
_RESERVED = (SUMMARY_SHEET, REFERENCE_SHEET, ISSUES_SHEET)
_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
_OUT_OF_SPEC_FILL = PatternFill("solid", fgColor="FFC7CE")
DEFAULT_CONFIG = Path("config.yaml")


class WideError(HrmConverterError):
    """The long workbook cannot be turned into wide tables."""


@dataclass
class WideResult:
    sheets: dict[str, pd.DataFrame]
    flagged: dict[str, list[tuple[int, str]]]  # sheet -> (row position, column) out of spec
    summary: pd.DataFrame
    reference: pd.DataFrame | None
    issues: list[Issue] = field(default_factory=list)

    @property
    def row_count(self) -> int:
        return sum(len(frame) for frame in self.sheets.values())

    @property
    def out_of_spec_count(self) -> int:
        return sum(len(cells) for cells in self.flagged.values())


@dataclass
class _Row:
    key: tuple[object, ...]
    values: dict[str, object] = field(default_factory=dict)
    limits: dict[str, Limits] = field(default_factory=dict)


@dataclass
class _Feature:
    rows: dict[tuple[object, ...], _Row] = field(default_factory=dict)
    units: dict[str, set[str]] = field(default_factory=dict)
    counts: dict[str, list[int]] = field(default_factory=dict)  # metric -> [values, limited, out]


def _number(value: object) -> float | None:
    if isinstance(value, bool) or not isinstance(value, int | float) or pd.isna(value):
        return None
    return float(value)


def _violations(value: object, limits: Limits) -> list[str]:
    number = _number(value)
    if number is None:
        return []
    found = []
    if limits.lsl is not None and number < limits.lsl:
        found.append("< LSL")
    if limits.usl is not None and number > limits.usl:
        found.append("> USL")
    return found


def _metric_order(metrics: Sequence[str], config: Config) -> list[str]:
    """Configured metrics first, in config order; anything else after, naturally sorted."""
    known: list[str] = []
    for names in config.metrics.positional.values():
        known.extend(names)
    known.extend(config.metrics.aliases.values())
    rank = {name: position for position, name in enumerate(dict.fromkeys(known))}
    return sorted(metrics, key=lambda m: (rank.get(m, len(rank)), natural_key(m)))


def _sheet_name(feature_type: str, taken: set[str]) -> str:
    base = _BAD_SHEET_CHARS.sub("_", feature_type).strip() or "Feature"
    name, counter = base[:31], 1
    while name.lower() in taken:
        counter += 1
        suffix = f"_{counter}"
        name = base[: 31 - len(suffix)] + suffix
    taken.add(name.lower())
    return name


def build_wide(
    long: pd.DataFrame, reference: Reference | None, config: Config, issues: IssueCollector
) -> WideResult:
    """Reshape the long table into one wide frame per feature type."""
    missing = [column for column in LONG_COLUMNS if column not in long.columns]
    if missing:
        raise WideError(f"The long table lacks the column(s): {', '.join(missing)}.")

    features: dict[str, _Feature] = {}
    repeats: dict[tuple[object, ...], int] = {}
    matcher = LimitMatcher(reference, config, issues) if reference is not None else None

    for record in long[list(LONG_COLUMNS)].itertuples(index=False, name=None):
        item = dict(zip(LONG_COLUMNS, record, strict=True))
        if is_blank(item["Value"]):
            continue
        feature_type, metric = str(item["Feature_Type"]), str(item["Metric"])
        feature = features.setdefault(feature_type, _Feature())

        base = (
            *(item[c] for c in META_COLUMNS),
            item["Feature_Number"],
            item["Unit"],
            item["Source_File"],
        )
        repeat_key = (feature_type, *base, metric)
        repeat = repeats.get(repeat_key, 0)
        repeats[repeat_key] = repeat + 1
        row = feature.rows.setdefault((*base, repeat), _Row((*base, repeat)))
        row.values[metric] = item["Value"]

        unit = "" if is_blank(item["Unit_of_Measurement"]) else str(item["Unit_of_Measurement"])
        feature.units.setdefault(metric, set()).add(unit)
        counts = feature.counts.setdefault(metric, [0, 0, 0])
        counts[0] += 1

        if matcher is not None:
            limits = matcher.limits(item)
            if not limits.is_empty():
                row.limits[metric] = limits
                counts[1] += 1
                counts[2] += bool(_violations(item["Value"], limits))

    taken = {name.lower() for name in _RESERVED}
    sheets: dict[str, pd.DataFrame] = {}
    flagged: dict[str, list[tuple[int, str]]] = {}
    summary_rows: list[tuple[object, ...]] = []
    for feature_type in sorted(features, key=natural_key):
        name = _sheet_name(feature_type, taken)
        sheets[name], flagged[name] = _feature_frame(feature_type, features[feature_type], config)
        for metric in _metric_order(list(features[feature_type].counts), config):
            summary_rows.append((feature_type, metric, *features[feature_type].counts[metric]))

    reference_frame = None
    if reference is not None:
        reference_frame = reference.frame.copy()
        reference_frame["Matched_Groups"] = [
            reference.matched.get(position + 2, 0) for position in range(len(reference_frame))
        ]
        used_rows = {rule.excel_row for rule in reference.rules}
        clashing = {row for conflict in reference.conflicts for row in conflict}
        unused = sorted(row for row in used_rows - clashing if reference.matched.get(row, 0) == 0)
        if unused:
            issues.warning(
                "reference_unused",
                f"{len(unused)} reference row(s) matched no measurement: row(s) "
                f"{', '.join(str(r) for r in unused)}. Check spelling of the key cells.",
                source_file=reference.source,
            )

    summary: pd.DataFrame = pd.DataFrame(summary_rows, columns=list(SUMMARY_COLUMNS), dtype=object)
    return WideResult(sheets, flagged, summary, reference_frame, issues.reportable())


def _feature_frame(
    feature_type: str, feature: _Feature, config: Config
) -> tuple[pd.DataFrame, list[tuple[int, str]]]:
    metrics = _metric_order(list(feature.counts), config)
    number_column = feature_type if feature_type not in LONG_COLUMNS else f"{feature_type}_Number"

    def header(metric: str) -> str:
        units = feature.units[metric] - {""}
        return f"{metric} ({next(iter(units))})" if len(units) == 1 else metric

    limit_kinds = {
        metric: [
            kind
            for kind in LIMIT_COLUMNS
            if any(
                getattr(row.limits[metric], kind.lower()) is not None
                for row in feature.rows.values()
                if metric in row.limits
            )
        ]
        for metric in metrics
    }
    any_limits = any(limit_kinds.values())
    any_repeat = any(row.key[-1] for row in feature.rows.values())

    columns = [*META_COLUMNS, "Unit", number_column]
    for metric in metrics:
        columns.append(header(metric))
        columns.extend(f"{metric}_{kind}" for kind in limit_kinds[metric])
    if any_limits:
        columns.append(OUT_OF_SPEC)
    if any_repeat:
        columns.append("Repeat")
    columns.append("Source_File")

    def sort_key(row: _Row) -> tuple[object, ...]:
        *meta, number, unit, source, repeat = row.key
        return tuple(natural_key(v) for v in (*meta, number, unit, source, repeat))

    records: list[dict[str, object]] = []
    flagged: list[tuple[int, str]] = []
    for position, row in enumerate(sorted(feature.rows.values(), key=sort_key)):
        *meta, number, unit, source, repeat = row.key
        record: dict[str, object] = dict(zip(META_COLUMNS, meta, strict=True))
        record.update({"Unit": unit, number_column: number, "Source_File": source})
        if any_repeat:
            record["Repeat"] = int(str(repeat)) + 1
        out_of_spec: list[str] = []
        for metric in metrics:
            record[header(metric)] = row.values.get(metric)
            limits = row.limits.get(metric, Limits())
            for kind in limit_kinds[metric]:
                record[f"{metric}_{kind}"] = getattr(limits, kind.lower())
            violations = _violations(row.values.get(metric), limits)
            if violations:
                flagged.append((position, header(metric)))
                out_of_spec.extend(f"{metric} {text}" for text in violations)
        if any_limits:
            record[OUT_OF_SPEC] = "; ".join(out_of_spec)
        records.append(record)

    frame: pd.DataFrame = pd.DataFrame(records, columns=columns, dtype=object)
    return frame, flagged


def read_long(path: Path, config: Config) -> pd.DataFrame:
    """Read the ``Long`` sheet of a converter output workbook."""
    try:
        frame: pd.DataFrame = pd.read_excel(
            path, sheet_name=config.output.sheet_name, dtype=object, keep_default_na=False
        )
    except Exception as exc:
        raise WideError(
            f"Long workbook cannot be read: {path} (sheet '{config.output.sheet_name}'; "
            f"{type(exc).__name__}: {exc})"
        ) from exc
    return frame


def _table_name(sheet: str, taken: set[str]) -> str:
    base = "tbl" + re.sub(r"\W", "_", sheet)
    name, counter = base, 1
    while name.lower() in taken:
        counter += 1
        name = f"{base}_{counter}"
    taken.add(name.lower())
    return name


def write_wide(result: WideResult, path: Path) -> Path:
    """Write the wide workbook: feature sheets, Limit_Summary, Reference, Wide_Issues."""
    if path.suffix.lower() != ".xlsx":
        raise OutputError(f"Wide output filename must end with .xlsx: {path.name}")
    frames = dict(result.sheets)
    frames[SUMMARY_SHEET] = result.summary
    if result.reference is not None:
        frames[REFERENCE_SHEET] = result.reference
    frames[ISSUES_SHEET] = issues_frame(result.issues)

    temporary = path.with_name(f".{path.stem}.tmp.xlsx")
    tables: set[str] = set()
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
            for name, frame in frames.items():
                frame.to_excel(writer, sheet_name=name, index=False)
                sheet = writer.sheets[name]
                format_sheet(sheet, frame, _table_name(name, tables))
                for position, column in result.flagged.get(name, []):
                    letter = get_column_letter(list(frame.columns).index(column) + 1)
                    sheet[f"{letter}{position + 2}"].fill = _OUT_OF_SPEC_FILL
        os.replace(temporary, path)
    except OSError as exc:
        temporary.unlink(missing_ok=True)
        raise OutputError(
            f"Wide workbook cannot be written: {path} ({exc}). "
            f"Close the file if it is open in Excel and run again."
        ) from exc
    return path


def write_reference_template(long: pd.DataFrame, path: Path, config: Config) -> Path:
    """Write a starter reference workbook; never overwrites an existing one."""
    if path.exists():
        raise OutputError(
            f"Reference file already exists and was not overwritten: {path}. "
            f"Choose another name or delete it yourself."
        )
    template = reference_template(long)
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        template.to_excel(writer, sheet_name=config.wide.reference_sheet, index=False)
        format_sheet(writer.sheets[config.wide.reference_sheet], template, "tblLimits")
    return path


def format_summary(result: WideResult, output: Path) -> str:
    limited = int(result.summary["With_Limits"].sum()) if len(result.summary) else 0
    values = int(result.summary["Values"].sum()) if len(result.summary) else 0
    warnings = sum(1 for i in result.issues if i.severity.value == "WARNING")
    errors = sum(1 for i in result.issues if i.severity.value == "ERROR")
    return "\n".join(
        [
            "",
            "HRM wide format finished",
            f"  Feature sheets      : {', '.join(result.sheets) or 'none'}",
            f"  Wide rows           : {result.row_count}",
            f"  Values with limits  : {limited} of {values}",
            f"  Values out of spec  : {result.out_of_spec_count}",
            f"  Warnings / errors   : {warnings} / {errors}",
            f"  Output workbook     : {output}",
            *(
                ["  See the 'Wide_Issues' sheet of the output workbook for details."]
                if warnings or errors
                else []
            ),
        ]
    )


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="hrm_converter.wide",
        description="Build the wide workbook (one sheet per feature, with limits) from the "
        "long workbook written by hrm_converter.",
    )
    parser.add_argument("--input", type=Path, help="Long workbook (default: the converter output).")
    parser.add_argument("--reference", type=Path, help="Reference workbook with LSL/Target/USL.")
    parser.add_argument("--output", type=Path, help="Wide workbook to write.")
    parser.add_argument("--config", type=Path, help="config.yaml (default: ./config.yaml).")
    parser.add_argument(
        "--make-reference",
        type=Path,
        metavar="FILE",
        help="Write a starter reference workbook from the long data and stop.",
    )
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        config_path = args.config
        if config_path is None and DEFAULT_CONFIG.is_file():
            config_path = DEFAULT_CONFIG
        config = load_config(config_path)
        long_path = args.input or config.output.directory / config.output.filename
        long = read_long(long_path, config)

        if args.make_reference is not None:
            created = write_reference_template(long, args.make_reference, config)
            print(f"Reference template written: {created.resolve()}")
            print("Fill in LSL / Target / USL, then run again with --reference.")
            return 0

        issues = IssueCollector()
        reference_path = args.reference or config.wide.reference_path
        output = args.output or long_path.with_name(config.wide.filename)
        # Issues go to the log file and the Wide_Issues sheet, not to the console.
        setup_logging(config.logging.directory, config.logging.level)
        try:
            reference = load_reference(reference_path, config, issues) if reference_path else None
            result = build_wide(long, reference, config, issues)
            write_wide(result, output)
        finally:
            close_logging()
    except HrmConverterError as exc:
        print(f"Stopped: {exc}", file=sys.stderr)
        return 1
    print(format_summary(result, output.resolve()))
    return 0


if __name__ == "__main__":
    sys.exit(main())
